"""
Generate sample resume files (PDF + DOCX) used by the test suite.

Dev-only helper -- needs `pip install -r requirements-dev.txt`.
Run:  python tests/make_fixtures.py
"""
import pathlib
import sys

FIXTURES = pathlib.Path(__file__).parent / "fixtures"

RESUME_LINES = [
    ("H", "ANANYA SHARMA"),
    ("P", "ananya.sharma@gndec.ac.in | +91 98765 43210"),
    ("P", "linkedin.com/in/ananya-sharma-dev | github.com/ananyasharma-dev"),
    ("", ""),
    ("S", "CAREER OBJECTIVE"),
    ("P", "Final-year Computer Science Engineering student seeking an AI/ML Engineer role. "
          "Experienced in building machine learning and retrieval-augmented systems in Python."),
    ("", ""),
    ("S", "EDUCATION"),
    ("B", "B.Tech in Computer Science Engineering, Guru Nanak Dev Engineering College, 2022 - 2026, CGPA: 8.6/10"),
    ("B", "Senior Secondary (Class XII), Delhi Public School, 2022, 92%"),
    ("", ""),
    ("S", "TECHNICAL SKILLS"),
    ("P", "Languages: Python, Java, JavaScript, SQL, C++"),
    ("P", "ML / AI: Machine Learning, Deep Learning, PyTorch, scikit-learn, NLP, Hugging Face, Embeddings"),
    ("P", "Backend: FastAPI, Flask, REST APIs, PostgreSQL, MongoDB"),
    ("P", "Tools: Git, Docker, Linux, Postman, Jupyter, Pandas, NumPy"),
    ("", ""),
    ("S", "EXPERIENCE"),
    ("B", "Machine Learning Intern, Infosys Springboard (Jun 2025 - Aug 2025)"),
    ("B", "Built a document classification pipeline in PyTorch that improved accuracy by 18% over the "
          "existing baseline across 12,000 records."),
    ("B", "Automated the data-cleaning workflow with Pandas, reducing manual preprocessing time by 6 hours per week."),
    ("B", "Deployed the trained model behind a FastAPI service containerised with Docker."),
    ("", ""),
    ("S", "PROJECTS"),
    ("B", "Semantic Research Assistant - Built a retrieval-augmented generation system over 500+ research papers "
          "using sentence embeddings and FAISS vector search, served through a FastAPI backend."),
    ("B", "Campus Placement Predictor - Trained scikit-learn classifiers on 3,000 student records; achieved "
          "87% accuracy with cross-validation and feature engineering."),
    ("B", "MediTrack Web App - Full-stack React and Node.js application with MongoDB, JWT authentication and "
          "a responsive Tailwind CSS interface used by 200 students."),
    ("", ""),
    ("S", "CERTIFICATIONS"),
    ("B", "Machine Learning Specialization - Coursera (Stanford Online), 2024"),
    ("B", "AWS Certified Cloud Practitioner, 2025"),
    ("", ""),
    ("S", "ACHIEVEMENTS"),
    ("B", "Winner, Smart India Hackathon internal round 2025, among 120 competing teams."),
    ("B", "Solved 450+ data structures and algorithms problems on LeetCode."),
    ("B", "Published a paper on NLP-based text summarisation at a national conference, 2025."),
]


def write_docx(path: pathlib.Path) -> None:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    document.styles["Normal"].font.size = Pt(10)

    for kind, text in RESUME_LINES:
        if kind == "" or not text:
            continue
        if kind == "H":
            para = document.add_paragraph()
            run = para.add_run(text)
            run.bold = True
            run.font.size = Pt(16)
        elif kind == "S":
            para = document.add_paragraph()
            run = para.add_run(text)
            run.bold = True
            run.font.size = Pt(12)
        elif kind == "B":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)

    document.save(path)


def write_pdf(path: pathlib.Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    width, height = A4
    left, right = 18 * mm, width - 18 * mm
    y = height - 20 * mm
    c = canvas.Canvas(str(path), pagesize=A4)

    def wrap(text, font, size, max_width):
        c.setFont(font, size)
        words, line, out = text.split(), "", []
        for word in words:
            trial = f"{line} {word}".strip()
            if c.stringWidth(trial, font, size) <= max_width:
                line = trial
            else:
                out.append(line)
                line = word
        if line:
            out.append(line)
        return out

    for kind, text in RESUME_LINES:
        if y < 25 * mm:
            c.showPage()
            y = height - 20 * mm
        if kind == "" or not text:
            y -= 4
            continue
        if kind == "H":
            c.setFont("Helvetica-Bold", 16)
            c.drawString(left, y, text)
            y -= 18
        elif kind == "S":
            c.setFont("Helvetica-Bold", 11)
            c.drawString(left, y, text)
            y -= 4
            c.line(left, y, right, y)
            y -= 12
        elif kind == "B":
            for i, ln in enumerate(wrap(text, "Helvetica", 9.5, right - left - 12)):
                c.setFont("Helvetica", 9.5)
                if i == 0:
                    c.drawString(left, y, "•")
                c.drawString(left + 12, y, ln)
                y -= 12
        else:
            for ln in wrap(text, "Helvetica", 9.5, right - left):
                c.setFont("Helvetica", 9.5)
                c.drawString(left, y, ln)
                y -= 12
    c.save()


def main() -> int:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    write_docx(FIXTURES / "sample_resume.docx")
    write_pdf(FIXTURES / "sample_resume.pdf")

    # A deliberately weak resume, to prove the ATS score actually discriminates.
    weak = FIXTURES / "weak_resume.docx"
    import docx
    d = docx.Document()
    d.add_paragraph("Rahul Verma")
    d.add_paragraph("rahul@example.com")
    d.add_paragraph("I am a student. I know computers and want a good job in a company.")
    d.add_paragraph("I have done some work with HTML and made a website.")
    d.save(weak)

    # Not a resume at all -- used for the invalid-input tests.
    (FIXTURES / "not_a_resume.txt").write_bytes(b"this is plain text, not a resume")
    (FIXTURES / "fake.pdf").write_bytes(b"NOT-REALLY-A-PDF" * 20)

    for f in sorted(FIXTURES.iterdir()):
        print(f"  {f.name:24} {f.stat().st_size:>8,} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
